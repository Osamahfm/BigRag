"""Bounded, non-executing extraction with page / section provenance."""
import csv
import hashlib
import io
import re
import zipfile
from pathlib import PurePosixPath

from rag.errors import RAGError
from rag.text import clean

SUPPORTED = {".pdf", ".txt", ".md", ".docx", ".csv", ".xlsx"}
MAX_EXTRACTED_CHARS = 5_000_000
MAX_UNITS = 10000


def document_hash(data):
    return hashlib.sha256(data).hexdigest()


def safe_filename(name):
    name = PurePosixPath(name.replace("\\", "/")).name
    name = re.sub(r'[\x00-\x1f<>:"/\\|?*\[\]]', "_", name).strip(" .")
    if not name or len(name) > 200:
        raise RAGError("Filename is empty or too long.")
    return name


def validate_upload(name, data, settings):
    name = safe_filename(name)
    ext = PurePosixPath(name).suffix.casefold()
    if ext not in SUPPORTED:
        raise RAGError("Unsupported format. Use PDF, TXT, MD, DOCX, CSV, or XLSX.")
    if not data or len(data) > settings.max_upload_mb * 1024 * 1024:
        raise RAGError(f"File must be nonempty and at most {settings.max_upload_mb} MB.")
    if ext == ".pdf" and not data.lstrip().startswith(b"%PDF-"):
        raise RAGError("File does not contain a valid PDF header.")
    if ext in {".docx", ".xlsx"}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                    raise RAGError("Office archive exceeds safe decompression limits.")
                expected = "word/document.xml" if ext == ".docx" else "xl/workbook.xml"
                if expected not in archive.namelist():
                    raise RAGError("Office file has an invalid internal structure.")
        except zipfile.BadZipFile:
            raise RAGError("Invalid Office archive.") from None
    return name, ext


def extract(name, data, settings):
    name, ext = validate_upload(name, data, settings)
    units, warnings = [], []
    used_chars = 0

    def add(text, page=None, section="Document"):
        nonlocal used_chars
        text = clean(text)
        used_chars += len(text)
        if used_chars > MAX_EXTRACTED_CHARS or len(units) >= MAX_UNITS:
            raise RAGError("Extracted content exceeds local processing limits; split this file.")
        if text:
            units.append({"text": text, "page": page, "section": str(section), "source": name})

    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted and not reader.decrypt(""):
                raise RAGError("Password protected PDF. Upload an unlocked copy.")
            pages = len(reader.pages)
            if pages > 2000:
                raise RAGError("PDF exceeds 2000 pages; split this file.")
            for number, page in enumerate(reader.pages, 1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                    warnings.append(f"Page {number} text extraction failed.")
                if not text.strip() and settings.ocr:
                    text = ocr_page(data, number - 1, warnings)
                if text.strip():
                    add(text, number, f"Page {number}")
                else:
                    warnings.append(f"Page {number} has no readable text; enable optional OCR for scans.")
        elif ext in {".txt", ".md", ".csv"}:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = data.decode("utf-16")
                except UnicodeDecodeError:
                    raise RAGError("Text files must be UTF-8 or UTF-16 encoded.") from None
            pages = None
            if ext == ".csv":
                reader = csv.reader(io.StringIO(text))
                header = next(reader, [])
                for number, row in enumerate(reader, 2):
                    add(" | ".join(f"{header[i] if i < len(header) else i+1}: {v}" for i, v in enumerate(row)),
                        section=f"Row {number}")
            elif ext == ".md":
                heading, lines = "Document", []
                for line in text.splitlines():
                    if re.match(r"^#{1,6}\s", line):
                        add("\n".join(lines), section=heading)
                        heading, lines = line.lstrip("# ").strip(), [line]
                    else:
                        lines.append(line)
                add("\n".join(lines), section=heading)
            else:
                add(text)
        elif ext == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc = Document(io.BytesIO(data))
            pages, heading, lines = None, "Document", []
            for block in doc.iter_inner_content():
                if isinstance(block, Paragraph):
                    if block.style and block.style.name.startswith("Heading"):
                        add("\n".join(lines), section=heading)
                        heading, lines = block.text, [block.text]
                    else:
                        lines.append(block.text)
                elif isinstance(block, Table):
                    lines.extend(" | ".join(c.text for c in row.cells) for row in block.rows)
            add("\n".join(lines), section=heading)
        else:
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            pages = None
            try:
                for sheet in book:
                    header = None
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        if number > 100000:
                            raise RAGError("Worksheet exceeds row limits; split this file.")
                        if not any(v is not None for v in row):
                            continue
                        if header is None:
                            header = [str(v) if v is not None else f"Column {i+1}" for i, v in enumerate(row)]
                            continue
                        add(" | ".join(f"{header[i] if i < len(header) else i+1}: {v}"
                                       for i, v in enumerate(row) if v is not None),
                            section=f"{sheet.title}, Row {number}")
            finally:
                book.close()
    except RAGError:
        raise
    except Exception:
        raise RAGError(f"{name}: unreadable file or missing format dependency.") from None
    if not units:
        raise RAGError(f"{name}: no readable text found. " + " ".join(warnings[:3]))
    return units, {"source": name, "format": ext[1:], "bytes": len(data), "pages": pages,
                   "units": len(units), "hash": document_hash(data)}, warnings


def ocr_page(data, page_number, warnings):
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            page = doc[page_number]
            if page.rect.width * page.rect.height > 4_000_000:
                raise ValueError("Page too large")
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            return pytesseract.image_to_string(Image.open(io.BytesIO(pix.tobytes("png"))),
                                              lang="eng+ara", timeout=30)
    except Exception:
        warnings.append(f"Page {page_number+1}: OCR unavailable; install optional packages, Tesseract and eng/ara language data.")
        return ""

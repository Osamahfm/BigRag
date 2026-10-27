import io
import zipfile
from dataclasses import replace
from unittest.mock import Mock, patch

import pytest
from conftest import make_pdf
from docx import Document
from openpyxl import Workbook

from rag.chunking import chunk_units
from rag.errors import RAGError
from rag.ingestion import extract, safe_filename


def test_pdf_provenance(settings):
    units, metadata, warnings = extract("../../lecture.pdf", make_pdf(["Page one", "", "Page three"]), settings)
    assert [u["page"] for u in units] == [1, 3]
    assert metadata["pages"] == 3 and metadata["source"] == "lecture.pdf" and warnings
    chunks = chunk_units(units, metadata["hash"], settings)
    assert {c["page"] for c in chunks} == {1, 3}
    assert all(c["parent_id"] and c["document_id"] == metadata["hash"] for c in chunks)


@pytest.mark.parametrize("raw", [b"not PDF", make_pdf([""]), make_pdf(["secret"], "password")])
def test_bad_pdf(raw, settings):
    with pytest.raises(RAGError):
        extract("bad.pdf", raw, settings)


def test_unreadable_page_retains_page_number(settings):
    reader = Mock(is_encrypted=False)
    bad, good = Mock(), Mock()
    bad.extract_text.side_effect = ValueError("unreadable")
    good.extract_text.return_value = "Good page"
    reader.pages = [bad, good]
    with patch("pypdf.PdfReader", return_value=reader):
        units, _, warnings = extract("pages.pdf", b"%PDF-1.7", settings)
    assert units[0]["page"] == 2 and warnings


@pytest.mark.parametrize("name,data,section", [("test.txt", "مرحبا بالعالم".encode(), "Document"),
    ("test.md", b"# Title\nhello world", "Title"),
    ("test.csv", b"Name,Score\nAlice,90\nBob,80", "Row 2")])
def test_text_formats(name, data, section, settings):
    units, _, _ = extract(name, data, settings)
    assert units[0]["section"] == section and units[0]["page"] is None


def test_docx_tables_and_headings(settings):
    doc = Document()
    doc.add_heading("Results", 1)
    doc.add_paragraph("Here are the results.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Accuracy", "95%"
    buffer = io.BytesIO()
    doc.save(buffer)
    units, _, _ = extract("result.docx", buffer.getvalue(), settings)
    assert units[0]["section"] == "Results"
    assert "Accuracy | 95%" in units[0]["text"]


def test_xlsx_table_extraction(settings):
    book = Workbook()
    sheet = book.active
    sheet.title = "Scores"
    sheet.append(["Name", "Score"])
    sheet.append(["Alice", 95])
    buffer = io.BytesIO()
    book.save(buffer)
    units, _, _ = extract("results.xlsx", buffer.getvalue(), settings)
    assert units[0]["section"] == "Scores, Row 2" and "Score: 95" in units[0]["text"]


@pytest.mark.parametrize("name,data", [("run.exe", b"code"), ("blank.txt", b""), ("bad.docx", b"notzip"),
                                      ("binary.txt", b"\xff\xff\xff")])
def test_upload_validation(name, data, settings):
    with pytest.raises(RAGError):
        extract(name, data, settings)


def test_size_limit_and_archive_limit(settings):
    with pytest.raises(RAGError, match="at most"):
        extract("large.txt", b"x"*(1024*1024+1), replace(settings, max_upload_mb=1))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x"*100_000_000)
        archive.writestr("extra", b"x"*10_000_000)
    with pytest.raises(RAGError, match="decompression"):
        extract("bomb.docx", buffer.getvalue(), settings)


def test_filename_sanitization():
    assert safe_filename(r"C:\secret\..\[unsafe].txt") == "_unsafe_.txt"
    with pytest.raises(RAGError):
        safe_filename("...")


@pytest.mark.parametrize("strategy", ["fixed", "recursive", "sentence", "parent-child", "semantic"])
def test_chunk_strategies(strategy, settings, service):
    units = [{"text": "A sentence about MapReduce. "*50, "page": 7, "section": "Page 7", "source": "x.pdf"}]
    chunks = chunk_units(units, "doc", replace(settings, chunk_strategy=strategy), service)
    assert chunks and all(len(c["text"]) <= settings.chunk_size for c in chunks)
    assert all(c["page"] == 7 and c["source"] == "x.pdf" for c in chunks)
    assert len({c["id"] for c in chunks}) == len(chunks)
    assert all(len(c["parent_text"]) <= settings.parent_size for c in chunks)


def test_missing_ocr_is_actionable(settings):
    with patch.dict("sys.modules", {"pymupdf": None}):
        with pytest.raises(RAGError, match="no readable"):
            extract("scan.pdf", make_pdf([""]), replace(settings, ocr=True))
